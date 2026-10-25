# -*- coding: utf-8 -*-
import os, re
from docx import Document
B = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
out = []
def W(s): out.append(str(s))

d = Document(os.path.join(B, "01_Manuscript.docx"))
txt = [p.text for p in d.paragraphs]
body = "\n".join(txt).split("References")[0]
W("段落 %d  表格 %d" % (len(txt), len(d.tables)))
for i, t in enumerate(d.tables, 1):
    W("  Table %d: %d 行, 列 %s" % (i, len(t.rows), {len(r.cells) for r in t.rows}))

refs = re.findall(r"^(\d+)\.\s", "\n".join(txt), flags=re.M)
cites = set()
for grp in re.findall(r"\[([\d,\u2013\- ]+)\]", body):
    for part in re.split(r",", grp):
        part = part.strip()
        if re.search(r"\d+\s*[\u2013-]\s*\d+", part):
            a, b = re.split(r"[\u2013-]", part); cites.update(range(int(a), int(b)+1))
        elif part.isdigit():
            cites.add(int(part))
W("参考文献 %d 条; 正文引用 %s" % (len(refs), sorted(cites)))
W("未被引用: %s; 越界引用: %s" % ([n for n in range(1, len(refs)+1) if n not in cites],
                                [n for n in cites if n > len(refs)]))

ph = [x for x in txt if re.search(r"\[[A-Za-z]", x)]
W("\n占位符 %d 处:" % len(ph))
for x in ph: W("  - " + x[:80])

bad = [x for x in txt if len(x) > 80 and not re.search(r"[\.\)\u201d]$", x.strip())]
W("\n可疑截断 %d 处" % len(bad))
for x in bad[:5]: W("  - " + x[:70])

W("\n关键数字出现次数（正文）:")
for k in ["531","223","42.0","1.33","1.22","1.14","1.09","0.97","0.001","0.007"]:
    W("  %-6s %d" % (k, body.count(k)))

# 与表核对
W("\n表中数字抽查:")
for t in d.tables:
    cells = [c.text.strip() for r in t.rows for c in r.cells]
    for k in ["0.573","0.534","1.22 (1.05","0.634"]:
        if any(k in c for c in cells):
            W("  '%s' 出现在表格中" % k); break

open(os.path.join(B, "verify.txt"), "w", encoding="utf-8").write("\n".join(out))
print("ok")
